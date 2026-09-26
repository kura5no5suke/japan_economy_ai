import io
import re
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook

from database import create_database, save_economic_data


# 内閣府
# 機械受注統計調査報告
INDEX_URL = (
    "https://www.esri.cao.go.jp/"
    "jp/stat/juchu/juchu.html"
)

SOURCE = "内閣府"
INDICATOR = "機械受注（船舶・電力を除く民需）"
UNIT = "100万円"

SHEET_NAME = "季調・月次"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 "
        "(Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 "
        "Chrome/120 Safari/537.36"
    )
}


def find_latest_excel_url():
    """
    内閣府の公式ページから
    「主要需要者別機械受注額」の
    長期系列Excelを自動検出する。
    """

    print(
        "内閣府の機械受注統計ページを"
        "確認します..."
    )

    response = requests.get(
        INDEX_URL,
        headers=HEADERS,
        timeout=60,
    )

    print(
        "内閣府 HTTPステータス:",
        response.status_code,
    )

    response.raise_for_status()

    pattern = re.compile(
        r'href=["\']'
        r'([^"\']*chouki-1\.xlsx)'
        r'["\']',
        re.IGNORECASE,
    )

    matches = pattern.findall(
        response.text
    )

    if not matches:
        raise RuntimeError(
            "内閣府ページから"
            "長期系列Excelを"
            "見つけられませんでした。"
        )

    candidates = []

    for href in matches:
        url = urljoin(
            INDEX_URL,
            href,
        )

        # 現行の年別フォルダにある
        # 長期系列を優先する
        if re.search(
            r"/juchu/20\d{2}/",
            url,
        ):
            candidates.append(url)

    if candidates:
        excel_url = candidates[0]
    else:
        excel_url = urljoin(
            INDEX_URL,
            matches[0],
        )

    print(
        "取得対象Excel:",
        excel_url,
    )

    return excel_url


def download_excel(excel_url):
    """
    内閣府からExcelを取得する。
    """

    print(
        "機械受注の長期系列Excelを"
        "取得します..."
    )

    response = requests.get(
        excel_url,
        headers=HEADERS,
        timeout=60,
    )

    print(
        "Excel HTTPステータス:",
        response.status_code,
    )

    response.raise_for_status()

    if not response.content:
        raise RuntimeError(
            "Excelを取得できませんでした。"
        )

    return response.content


def normalize_year(value):
    """
    年を整数へ変換する。
    """

    if value is None:
        return None

    if isinstance(
        value,
        (int, float),
    ):
        year = int(value)

        if 2000 <= year <= 2100:
            return year

        return None

    text = str(value).strip()

    match = re.search(
        r"(20\d{2})",
        text,
    )

    if match:
        return int(
            match.group(1)
        )

    return None


def normalize_month(value):
    """
    月を整数へ変換する。
    """

    if value is None:
        return None

    if isinstance(
        value,
        (int, float),
    ):
        month = int(value)

        if 1 <= month <= 12:
            return month

        return None

    text = str(value).strip()

    match = re.search(
        r"(\d{1,2})",
        text,
    )

    if not match:
        return None

    month = int(
        match.group(1)
    )

    if 1 <= month <= 12:
        return month

    return None


def find_target_column(sheet):
    """
    「民需（船舶・電力を除く）」の
    データ列をヘッダーから探す。

    アップロードされた公式Excelでは
    J列が対象。
    """

    # まず現在の公式ExcelのJ列を
    # ヘッダー内容で確認する
    for row in range(
        1,
        min(sheet.max_row, 20) + 1,
    ):
        value = sheet.cell(
            row=row,
            column=10,
        ).value

        if value is None:
            continue

        text = (
            str(value)
            .replace(" ", "")
            .replace("　", "")
        )

        if (
            "船舶" in text
            and "電力" in text
        ):
            return 10

    # J列で確認できなかった場合は
    # ヘッダー全体から候補を探す
    for column in range(
        1,
        sheet.max_column + 1,
    ):
        combined_text = ""

        for row in range(
            1,
            min(sheet.max_row, 20) + 1,
        ):
            value = sheet.cell(
                row=row,
                column=column,
            ).value

            if value is not None:
                combined_text += (
                    str(value)
                    .replace(" ", "")
                    .replace("　", "")
                )

        if (
            "船舶" in combined_text
            and "電力" in combined_text
            and "民需" in combined_text
        ):
            return column

    raise RuntimeError(
        "「民需（船舶・電力を除く）」"
        "の列を見つけられませんでした。"
    )


def get_machinery_orders_data():
    """
    最新Excelから
    季節調整済み月次の
    「民需（船舶・電力を除く）」
    を取得する。
    """

    excel_url = (
        find_latest_excel_url()
    )

    excel_data = download_excel(
        excel_url
    )

    workbook = load_workbook(
        filename=io.BytesIO(
            excel_data
        ),
        read_only=True,
        data_only=True,
    )

    if SHEET_NAME not in (
        workbook.sheetnames
    ):
        workbook.close()

        raise RuntimeError(
            f"Excelに「{SHEET_NAME}」"
            "シートがありません。"
        )

    sheet = workbook[
        SHEET_NAME
    ]

    target_column = (
        find_target_column(sheet)
    )

    records = []

    current_year = None

    for row in range(
        1,
        sheet.max_row + 1,
    ):
        year_value = sheet.cell(
            row=row,
            column=2,
        ).value

        month_value = sheet.cell(
            row=row,
            column=3,
        ).value

        value = sheet.cell(
            row=row,
            column=target_column,
        ).value

        detected_year = (
            normalize_year(
                year_value
            )
        )

        if detected_year is not None:
            current_year = (
                detected_year
            )

        month = normalize_month(
            month_value
        )

        if (
            current_year is None
            or month is None
            or value is None
        ):
            continue

        try:
            numeric_value = float(
                value
            )
        except (
            TypeError,
            ValueError,
        ):
            continue

        date = (
            f"{current_year:04d}-"
            f"{month:02d}"
        )

        records.append(
            (
                date,
                numeric_value,
            )
        )

    workbook.close()

    # 同じ年月が存在した場合に備えて
    # 重複を除去する
    records = sorted(
        dict(records).items(),
        key=lambda x: x[0],
    )

    return records


def save_machinery_orders_data(
    records,
):
    """
    database.py の共通関数を使って
    SQLiteへ保存する。
    """

    saved_count = 0

    for date, value in records:
        save_economic_data(
            SOURCE,
            INDICATOR,
            date,
            value,
            UNIT,
        )

        saved_count += 1

    return saved_count


def calculate_change(
    previous_value,
    latest_value,
):
    """
    前月比を計算する。
    """

    if previous_value == 0:
        return None

    return (
        (
            latest_value
            - previous_value
        )
        / previous_value
        * 100
    )


def main():
    print()
    print(
        "===== 機械受注 ====="
    )

    create_database()

    records = (
        get_machinery_orders_data()
    )

    if not records:
        print(
            "保存できる機械受注データが"
            "ありません。"
        )
        return

    saved_count = (
        save_machinery_orders_data(
            records
        )
    )

    latest_date, latest_value = (
        records[-1]
    )

    if len(records) >= 2:
        (
            previous_date,
            previous_value,
        ) = records[-2]
    else:
        previous_date = None
        previous_value = None

    print()
    print(
        "===== 機械受注 取得結果 ====="
    )

    print(
        "保存件数:",
        saved_count,
    )

    if previous_date is not None:
        print(
            "前回:",
            previous_date,
            previous_value,
        )

    print(
        "最新年月:",
        latest_date,
    )

    print(
        "最新値:",
        latest_value,
        UNIT,
    )

    if previous_value is not None:
        change = calculate_change(
            previous_value,
            latest_value,
        )

        if change is not None:
            print(
                "前月比:",
                f"{change:+.2f}%",
            )


if __name__ == "__main__":
    main()