import io
import re
import requests
from openpyxl import load_workbook
from database import save_economic_data


EXCEL_URL = (
    "https://www.e-stat.go.jp/stat-search/file-download"
    "?statInfId=000031831358&fileKind=0"
)

SOURCE = "e-Stat"
INDICATOR = "完全失業率"
UNIT = "%"
SHEET_NAME = "季節調整値"


def parse_month(value):
    if value is None:
        return None

    match = re.search(r"(\d{1,2})\s*月", str(value))
    if not match:
        return None

    month = int(match.group(1))

    if 1 <= month <= 12:
        return month

    return None


def get_unemployment():
    response = requests.get(
        EXCEL_URL,
        headers={"User-Agent": "Mozilla/5.0"},
        timeout=60,
    )
    response.raise_for_status()

    print("HTTPステータス:", response.status_code)

    workbook = load_workbook(
        io.BytesIO(response.content),
        read_only=True,
        data_only=True,
    )

    sheet = workbook[SHEET_NAME]

    # T列 = 完全失業率（季節調整値・男女計）
    unemployment_column = 20

    current_year = None
    pending_january_year = None
    records = []

    for row in range(11, sheet.max_row + 1):
        year_cell = sheet.cell(row, 1).value
        month = parse_month(sheet.cell(row, 2).value)
        value = sheet.cell(row, unemployment_column).value

        if month is None or value in (None, ""):
            continue

        # 各年の1月行には「令和 8年」のような和暦が入る。
        # 次の2月行には西暦（2026）が入るため、
        # 1月の値を一時保持し、西暦を確認してから保存する。
        if month == 1:
            pending_january_year = (row, value)
            continue

        if isinstance(year_cell, (int, float)):
            year = int(year_cell)

            if 1900 <= year <= 2100:
                current_year = year

                if pending_january_year is not None:
                    _, january_value = pending_january_year
                    if current_year >= 2000:
                        records.append(
                            (f"{current_year}-01", float(january_value))
                        )
                    pending_january_year = None

        if current_year is None:
            continue

        if current_year < 2000:
            continue

        records.append(
            (f"{current_year}-{month:02d}", float(value))
        )

    workbook.close()

    if not records:
        raise RuntimeError("完全失業率のデータを取得できませんでした")

    # 同じ年月が存在した場合は後から読んだ値を採用
    records = sorted(dict(records).items())

    print()
    print("===== 完全失業率（季節調整値） =====")

    count = 0

    for date, value in records:
        print(
            "日付:", date,
            "| 完全失業率:", value, UNIT
        )

        save_economic_data(
            source=SOURCE,
            indicator=INDICATOR,
            date=date,
            value=value,
            unit=UNIT,
        )

        count += 1

    print()
    print("保存件数:", count)

    previous_date, previous_value = records[-2]
    latest_date, latest_value = records[-1]

    print("前回:", previous_date, previous_value, UNIT)
    print("最新:", latest_date, latest_value, UNIT)


if __name__ == "__main__":
    get_unemployment()
