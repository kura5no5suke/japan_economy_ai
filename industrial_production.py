import io

import requests
from openpyxl import load_workbook

from database import create_database, save_economic_data


# e-Stat
# 鉱工業生産・出荷・在庫指数
# 2020年基準
# 業種別 / 月次 / 季節調整済指数
EXCEL_URL = (
    "https://www.e-stat.go.jp/stat-search/file-download"
    "?statInfId=000040172363&fileKind=0"
)

SOURCE = "e-Stat"
INDICATOR = "鉱工業生産指数"
UNIT = "2020年=100"

SHEET_NAME = "生産"
TARGET_NAME = "鉱工業"


def download_excel():
    print("e-Statから最新Excelを取得します...")

    response = requests.get(
        EXCEL_URL,
        timeout=60,
    )

    print(
        "e-Stat HTTPステータス:",
        response.status_code,
    )

    response.raise_for_status()

    if not response.content:
        raise RuntimeError(
            "e-StatからExcelを取得できませんでした。"
        )

    return response.content


def normalize_month(value):
    """
    Excelの年月を YYYY-MM に変換する。

    例:
    202607 -> 2026-07
    """

    if value is None:
        return None

    # 数値として入っている場合
    if isinstance(value, (int, float)):
        text = str(int(value))
    else:
        text = str(value).strip()

    # 202607.0 のような形式への対策
    if text.endswith(".0"):
        text = text[:-2]

    if len(text) != 6:
        return None

    if not text.isdigit():
        return None

    year = text[:4]
    month = text[4:6]

    try:
        month_number = int(month)
    except ValueError:
        return None

    if not 1 <= month_number <= 12:
        return None

    return f"{year}-{month}"


def get_industrial_production_data():
    excel_data = download_excel()

    workbook = load_workbook(
        filename=io.BytesIO(excel_data),
        read_only=True,
        data_only=True,
    )

    if SHEET_NAME not in workbook.sheetnames:
        raise RuntimeError(
            f"Excelに「{SHEET_NAME}」シートがありません。"
        )

    sheet = workbook[SHEET_NAME]

    rows = list(
        sheet.iter_rows(values_only=True)
    )

    if not rows:
        raise RuntimeError(
            "Excelの生産シートが空です。"
        )

    # -------------------------
    # 鉱工業の行を探す
    # -------------------------

    target_row_index = None
    target_col_index = None

    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            if value is None:
                continue

            if str(value).strip() == TARGET_NAME:
                target_row_index = row_index
                target_col_index = col_index
                break

        if target_row_index is not None:
            break

    if target_row_index is None:
        raise RuntimeError(
            "Excelから「鉱工業」の行を"
            "見つけられませんでした。"
        )

    # -------------------------
    # 年月が入っている行を探す
    # -------------------------

    month_row_index = None

    for row_index in range(
        target_row_index - 1,
        -1,
        -1,
    ):
        valid_months = 0

        for value in rows[row_index]:
            if normalize_month(value):
                valid_months += 1

        if valid_months >= 2:
            month_row_index = row_index
            break

    if month_row_index is None:
        raise RuntimeError(
            "Excelから年月の行を"
            "見つけられませんでした。"
        )

    month_row = rows[month_row_index]
    target_row = rows[target_row_index]

    records = []

    column_count = min(
        len(month_row),
        len(target_row),
    )

    for col_index in range(column_count):
        # 「鉱工業」という文字そのものより
        # 左側はデータ列ではないため除外
        if col_index <= target_col_index:
            continue

        date = normalize_month(
            month_row[col_index]
        )

        if not date:
            continue

        value = target_row[col_index]

        try:
            numeric_value = float(value)
        except (TypeError, ValueError):
            continue

        records.append(
            (date, numeric_value)
        )

    workbook.close()

    # 同じ年月があった場合にも備えて重複除去
    records = sorted(
        dict(records).items(),
        key=lambda x: x[0],
    )

    return records


def save_industrial_production_data(records):
    saved_count = 0

    for date, value in records:
        save_economic_data(
            SOURCE,
            INDICATOR,
            date,
            value,
            UNIT,
        )

        saved_count += 1

    return saved_count


def main():
    print()
    print(
        "===== 鉱工業生産指数 ====="
    )

    create_database()

    records = (
        get_industrial_production_data()
    )

    if not records:
        print(
            "保存できる鉱工業生産指数が"
            "ありません。"
        )
        return

    saved_count = (
        save_industrial_production_data(
            records
        )
    )

    latest_date, latest_value = records[-1]

    if len(records) >= 2:
        previous_date, previous_value = (
            records[-2]
        )
    else:
        previous_date = None
        previous_value = None

    print()
    print(
        "===== 鉱工業生産指数 取得結果 ====="
    )

    print(
        "保存件数:",
        saved_count,
    )

    if previous_date is not None:
        print(
            "前回:",
            previous_date,
            previous_value,
        )

    print(
        "最新年月:",
        latest_date,
    )

    print(
        "最新指数:",
        latest_value,
        "(2020年=100)",
    )


if __name__ == "__main__":
    main()