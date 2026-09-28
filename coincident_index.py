import io
import re
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook

from database import create_database, save_economic_data


# 内閣府
# 景気動向指数
INDEX_URL = (
    "https://www.esri.cao.go.jp/"
    "jp/stat/di/di.html"
)

SOURCE = "内閣府"
INDICATOR = "景気動向指数（CI一致指数）"
UNIT = "2020年=100"

SHEET_NAME = "指数 Indexes"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 "
        "(Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 "
        "Chrome/120 Safari/537.36"
    )
}


def find_latest_excel_url():
    """
    内閣府の景気動向指数ページから
    「長期系列（CI指数、DI指数、DI累積指数）」
    のExcelを自動検出する。
    """

    print(
        "内閣府の景気動向指数ページを"
        "確認します..."
    )

    response = requests.get(
        INDEX_URL,
        headers=HEADERS,
        timeout=60,
    )

    print(
        "内閣府 HTTPステータス:",
        response.status_code,
    )

    response.raise_for_status()

    # 長期系列のExcelは
    # 0907ci.xlsx のような形式。
    #
    # ci1 / ci2 / ci3 や
    # ci_cont1 などの個別系列を
    # 誤って取得しないように、
    # 数字4桁 + ci.xlsx のみを対象にする。
    pattern = re.compile(
        r'href=["\']'
        r'([^"\']*?(\d{4})ci\.xlsx)'
        r'["\']',
        re.IGNORECASE,
    )

    matches = pattern.findall(
        response.text
    )

    if not matches:
        raise RuntimeError(
            "内閣府ページから"
            "景気動向指数の長期系列Excelを"
            "見つけられませんでした。"
        )

    candidates = []

    for href, date_code in matches:
        candidates.append(
            (
                int(date_code),
                urljoin(
                    INDEX_URL,
                    href,
                ),
            )
        )

    # ページ内に複数候補がある場合は、
    # MMDD形式のファイル名部分が
    # 最大のものを優先する。
    candidates.sort(
        key=lambda x: x[0],
        reverse=True,
    )

    excel_url = candidates[0][1]

    print(
        "取得対象Excel:",
        excel_url,
    )

    return excel_url


def download_excel(excel_url):
    """
    内閣府から長期系列Excelを取得する。
    """

    print(
        "景気動向指数の長期系列Excelを"
        "取得します..."
    )

    response = requests.get(
        excel_url,
        headers=HEADERS,
        timeout=60,
    )

    print(
        "Excel HTTPステータス:",
        response.status_code,
    )

    response.raise_for_status()

    if not response.content:
        raise RuntimeError(
            "Excelを取得できませんでした。"
        )

    return response.content


def normalize_year(value):
    """
    年を整数へ変換する。
    """

    if value is None:
        return None

    if isinstance(
        value,
        (int, float),
    ):
        year = int(value)

        if 1900 <= year <= 2100:
            return year

        return None

    text = str(value).strip()

    match = re.search(
        r"((?:19|20)\d{2})",
        text,
    )

    if match:
        return int(
            match.group(1)
        )

    return None


def normalize_month(value):
    """
    月を整数へ変換する。
    """

    if value is None:
        return None

    if isinstance(
        value,
        (int, float),
    ):
        month = int(value)

        if 1 <= month <= 12:
            return month

        return None

    text = str(value).strip()

    match = re.search(
        r"(\d{1,2})",
        text,
    )

    if not match:
        return None

    month = int(
        match.group(1)
    )

    if 1 <= month <= 12:
        return month

    return None


def find_coincident_column(sheet):
    """
    通常のCI一致指数の列を
    ヘッダーから確認する。

    現在の公式ExcelではE列。
    参考系列の一致指数ではなく、
    標準のCI一致指数を対象とする。
    """

    # 現行Excelでは
    # D～F列が通常のCI指数で、
    # E列が一致指数。
    combined_text = ""

    for row in range(
        1,
        min(sheet.max_row, 10) + 1,
    ):
        value = sheet.cell(
            row=row,
            column=5,
        ).value

        if value is not None:
            combined_text += (
                str(value)
                .replace(" ", "")
                .replace("　", "")
            )

    if (
        "一致指数" in combined_text
        or "Coincident" in combined_text
    ):
        return 5

    # E列で確認できなかった場合は、
    # ヘッダー全体から一致指数を探す。
    # ただし参考系列は除外する。
    for column in range(
        1,
        sheet.max_column + 1,
    ):
        texts = []

        for row in range(
            1,
            min(sheet.max_row, 10) + 1,
        ):
            value = sheet.cell(
                row=row,
                column=column,
            ).value

            if value is not None:
                texts.append(
                    str(value)
                    .replace(" ", "")
                    .replace("　", "")
                )

        combined = "".join(texts)

        if (
            (
                "一致指数" in combined
                or "Coincident" in combined
            )
            and "参考" not in combined
            and "Reference" not in combined
        ):
            return column

    raise RuntimeError(
        "CI一致指数の列を"
        "見つけられませんでした。"
    )


def get_coincident_index_data():
    """
    最新の長期系列Excelから
    CI一致指数を取得する。
    """

    excel_url = (
        find_latest_excel_url()
    )

    excel_data = download_excel(
        excel_url
    )

    workbook = load_workbook(
        filename=io.BytesIO(
            excel_data
        ),
        read_only=True,
        data_only=True,
    )

    if SHEET_NAME not in (
        workbook.sheetnames
    ):
        workbook.close()

        raise RuntimeError(
            f"Excelに「{SHEET_NAME}」"
            "シートがありません。"
        )

    sheet = workbook[
        SHEET_NAME
    ]

    target_column = (
        find_coincident_column(
            sheet
        )
    )

    records = []

    for row in range(
        1,
        sheet.max_row + 1,
    ):
        year_value = sheet.cell(
            row=row,
            column=2,
        ).value

        month_value = sheet.cell(
            row=row,
            column=3,
        ).value

        value = sheet.cell(
            row=row,
            column=target_column,
        ).value

        year = normalize_year(
            year_value
        )

        month = normalize_month(
            month_value
        )

        if (
            year is None
            or month is None
            or value is None
        ):
            continue

        try:
            numeric_value = float(
                value
            )
        except (
            TypeError,
            ValueError,
        ):
            continue

        date = (
            f"{year:04d}-"
            f"{month:02d}"
        )

        records.append(
            (
                date,
                numeric_value,
            )
        )

    workbook.close()

    # 同じ年月が存在した場合に備えて
    # 重複を除去して年月順に並べる。
    records = sorted(
        dict(records).items(),
        key=lambda x: x[0],
    )

    return records


def save_coincident_index_data(
    records,
):
    """
    database.py の共通関数を使って
    SQLiteへ保存する。
    """

    saved_count = 0

    for date, value in records:
        save_economic_data(
            SOURCE,
            INDICATOR,
            date,
            value,
            UNIT,
        )

        saved_count += 1

    return saved_count


def calculate_change(
    previous_value,
    latest_value,
):
    """
    前月比を計算する。
    """

    if previous_value == 0:
        return None

    return (
        (
            latest_value
            - previous_value
        )
        / abs(previous_value)
        * 100
    )


def main():
    print()
    print(
        "===== 景気動向指数 ====="
    )

    create_database()

    records = (
        get_coincident_index_data()
    )

    if not records:
        print(
            "保存できる景気動向指数データが"
            "ありません。"
        )
        return

    saved_count = (
        save_coincident_index_data(
            records
        )
    )

    latest_date, latest_value = (
        records[-1]
    )

    if len(records) >= 2:
        (
            previous_date,
            previous_value,
        ) = records[-2]
    else:
        previous_date = None
        previous_value = None

    print()
    print(
        "===== 景気動向指数 取得結果 ====="
    )

    print(
        "指標:",
        "CI一致指数",
    )

    print(
        "保存件数:",
        saved_count,
    )

    if previous_date is not None:
        print(
            "前回:",
            previous_date,
            previous_value,
        )

    print(
        "最新年月:",
        latest_date,
    )

    print(
        "最新値:",
        latest_value,
        UNIT,
    )

    if previous_value is not None:
        change = calculate_change(
            previous_value,
            latest_value,
        )

        if change is not None:
            print(
                "前月比:",
                f"{change:+.2f}%",
            )


if __name__ == "__main__":
    main()